"""
Step 12b: Build a SECOND, dialect-balanced validation batch.

The first validation_sample.xlsx was opportunistically filled (1-4 clips per
dialect). To report a per-dialect silver-label quality table we need even
coverage, so this samples `--per` clips from EACH dialect (default 2 -> ~30
clips across 15 dialects), preferring 1 code-switched + the rest Bengali-only
where CS clips exist. It EXCLUDES every clip already present in
validation_sample.xlsx so there is no overlap with the 31 you already filled.

Audio is copied into validation_clips_batch2/ and zipped so the clips can be
transcribed offline. Fill the human_transcript column, then score BOTH batches
together with 13_validation_wer.py (point it at the combined file).

Output (evaluation/results/):
  validation_sample_batch2.xlsx / .csv   <- fill human_transcript
  validation_clips_batch2/<dialect>/<clip>.wav + validation_clips_batch2.zip

Usage:
    python 12b_make_validation_batch2.py            # 2 per dialect
    python 12b_make_validation_batch2.py --per 3    # 3 per dialect
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import shutil
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(PROJECT_ROOT))
from fine_tuning.config import RESULTS_DIR  # noqa: E402

SEGMENTS = PROJECT_ROOT / "dataset_pipeline" / "segments"
COLS = ["clip_id", "dialect", "is_code_switched", "audio_path",
        "silver_transcript", "human_transcript"]


def _is_cs(row: dict) -> bool:
    return str(row.get("is_code_switched", "")).strip().lower() in ("true", "1")


def _used_clip_ids() -> set:
    """Clip ids already in validation_sample.xlsx/.csv (exclude to avoid overlap)."""
    used = set()
    xlsx = RESULTS_DIR / "validation_sample.xlsx"
    csvp = RESULTS_DIR / "validation_sample.csv"
    if xlsx.exists():
        from openpyxl import load_workbook
        ws = load_workbook(xlsx, read_only=True).active
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for r in ws.iter_rows(min_row=2, values_only=True):
            d = {hdr[i]: r[i] for i in range(len(hdr))}
            if d.get("clip_id"):
                used.add(str(d["clip_id"]))
    elif csvp.exists():
        with open(csvp, encoding="utf-8-sig") as f:
            used = {str(r["clip_id"]) for r in csv.DictReader(f) if r.get("clip_id")}
    return used


def _resolve_audio(clip_id: str, dialect: str) -> Path | None:
    p = SEGMENTS / str(dialect) / f"{clip_id}.wav"
    return p if p.exists() else None


def _write_xlsx(path: Path, records: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "validation"
    ws.append(COLS)
    for cell in ws[1]:
        cell.font = Font(name="Nirmala UI", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E7D32")
    for rec in records:
        ws.append([rec[c] for c in COLS])
    widths = {"clip_id": 22, "dialect": 12, "is_code_switched": 14,
              "audio_path": 52, "silver_transcript": 50, "human_transcript": 50}
    for i, col in enumerate(COLS, 1):
        ws.column_dimensions[chr(64 + i)].width = widths.get(col, 18)
    beng = Font(name="Nirmala UI", size=11)
    yellow = PatternFill("solid", fgColor="FFF9C4")
    for row in range(2, len(records) + 2):
        ws.cell(row, 5).font = beng
        ws.cell(row, 5).alignment = Alignment(wrap_text=True)
        ws.cell(row, 6).font = beng
        ws.cell(row, 6).fill = yellow
        ws.cell(row, 6).alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(path)


def main(per: int = 2, seed: int = 7):
    rng = random.Random(seed)
    test = json.load(open(RESULTS_DIR / "test_set.json", encoding="utf-8"))
    used = _used_clip_ids()
    print(f"[info] excluding {len(used)} clips already in validation_sample")

    by_dialect = defaultdict(list)
    for r in test:
        if str(r["clip_id"]) in used:
            continue
        by_dialect[r["dialect"]].append(r)

    sample = []
    for dialect in sorted(by_dialect):
        clips = by_dialect[dialect]
        cs = [c for c in clips if _is_cs(c)]
        bn = [c for c in clips if not _is_cs(c)]
        rng.shuffle(cs); rng.shuffle(bn)
        # prefer 1 CS (if any) + the rest BN
        pick = (cs[:1] + bn) if cs else bn
        chosen = 0
        for c in pick:
            if chosen >= per:
                break
            audio = _resolve_audio(c["clip_id"], dialect)
            if audio is None:        # only keep clips whose wav we can ship
                continue
            sample.append((c, audio))
            chosen += 1
        if chosen < per:
            print(f"  [warn] {dialect}: only {chosen}/{per} clips had resolvable audio")

    rng.shuffle(sample)
    records = [{
        "clip_id": c["clip_id"], "dialect": c.get("dialect", "unknown"),
        "is_code_switched": _is_cs(c),
        "audio_path": f"validation_clips_batch2/{c.get('dialect')}/{c['clip_id']}.wav",
        "silver_transcript": str(c.get("reference", "")).strip(),
        "human_transcript": "",
    } for c, _ in sample]

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    # csv + xlsx
    with open(RESULTS_DIR / "validation_sample_batch2.csv", "w",
              encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(records)
    _write_xlsx(RESULTS_DIR / "validation_sample_batch2.xlsx", records)

    # copy audio + zip
    clip_dir = RESULTS_DIR / "validation_clips_batch2"
    if clip_dir.exists():
        shutil.rmtree(clip_dir)
    for c, audio in sample:
        dst = clip_dir / str(c.get("dialect")) / f"{c['clip_id']}.wav"
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(audio, dst)
    zip_path = RESULTS_DIR / "validation_clips_batch2.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for wav in clip_dir.rglob("*.wav"):
            z.write(wav, wav.relative_to(RESULTS_DIR))

    n_cs = sum(1 for r in records if r["is_code_switched"])
    print(f"\nWrote {len(records)} clips across {len({r['dialect'] for r in records})} dialects")
    print(f"  CS: {n_cs} | BN: {len(records) - n_cs}")
    print(f"  -> {RESULTS_DIR / 'validation_sample_batch2.xlsx'}")
    print(f"  -> {zip_path}  (audio to listen to)")
    by = defaultdict(int)
    for r in records:
        by[r["dialect"]] += 1
    print("  per-dialect:", dict(sorted(by.items())))
    print("\nFill 'human_transcript' by listening, then score with 13_validation_wer.py")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--per", type=int, default=2, help="clips per dialect (default 2)")
    ap.add_argument("--seed", type=int, default=7)
    main(per=ap.parse_args().per, seed=ap.parse_args().seed)
