"""
Step 14: Score every model against HUMAN references on the validated test subset.

Step 13 reports silver-vs-human agreement (label quality). This script does the
complementary, reviewer-critical thing: it evaluates each ASR model's WER/CER
against the *human* transcripts on the human-validated test clips, giving a
gold-reference anchor for the benchmark (not silver-vs-silver). It also reports
each model's WER against the silver labels on the SAME clips, so the gap between
"human-referenced" and "silver-referenced" error is explicit.

It reads all filled validation_sample*.xlsx/.csv files (the union of every
validation batch) and every evaluation/results/*_predictions.json.

Convention: human transcript = reference (ground truth); model hypothesis from
the predictions JSON = hypothesis.

Usage:
    python 14_human_test_wer.py
"""
from __future__ import annotations

import csv
import glob
import json
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(PROJECT_ROOT))
from fine_tuning.config import RESULTS_DIR  # noqa: E402
from dataset_pipeline.text_normalization import full_normalize  # noqa: E402

# Friendly model labels (key = predictions-file stem without _predictions)
LABELS = {
    "whisper_small_zs":   "Whisper-small (zero-shot)",
    "whisper_large_v3":   "Whisper large-v3 (zero-shot)",
    "mms_bn":             "MMS-1B (zero-shot)",
    "wav2vec2_bn":        "wav2vec2-XLS-R-BN (zero-shot)",
    "ft_wav2vec2_all":    "wav2vec2 (fine-tuned)",
    "ft_whisper_bn_only": "Whisper BN_ONLY",
    "ft_whisper_cs_only": "Whisper CS_ONLY",
    "ft_whisper_all":     "Whisper ALL (proposed)",
}


def _read_one(stem: str) -> list[dict]:
    xlsx = RESULTS_DIR / f"{stem}.xlsx"
    csvp = RESULTS_DIR / f"{stem}.csv"
    if xlsx.exists():
        from openpyxl import load_workbook
        ws = load_workbook(xlsx, read_only=True).active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        return [{header[i]: ("" if r[i] is None else str(r[i]))
                 for i in range(len(header))}
                for r in ws.iter_rows(min_row=2, values_only=True)]
    if csvp.exists():
        with open(csvp, encoding="utf-8-sig") as f:
            return [dict(r) for r in csv.DictReader(f)]
    return []


def _human_refs() -> dict:
    """clip_id -> human_transcript over all filled validation batches."""
    stems = {Path(p).stem for p in
             glob.glob(str(RESULTS_DIR / "validation_sample*.xlsx"))
             + glob.glob(str(RESULTS_DIR / "validation_sample*.csv"))
             if not Path(p).name.startswith("~$")}
    refs = {}
    for stem in sorted(stems):
        for r in _read_one(stem):
            cid = str(r.get("clip_id", ""))
            h = str(r.get("human_transcript", "")).strip()
            if cid and h:
                refs.setdefault(cid, h)
    return refs


def _score(refs: list[str], hyps: list[str], metric: str):
    from jiwer import cer as jcer, wer as jwer
    fn = jwer if metric == "wer" else jcer
    pairs = [(full_normalize(r), full_normalize(h))
             for r, h in zip(refs, hyps) if full_normalize(r).strip()]
    if not pairs:
        return None
    r, h = zip(*pairs)
    return round(fn(list(r), list(h)), 4)


def main():
    human = _human_refs()
    if not human:
        print("[ERROR] No filled human_transcript found in any validation_sample*.\n"
              "        Fill a validation sheet (see 12b_make_validation_batch2.py) first.")
        raise SystemExit(1)
    print(f"[info] {len(human)} human-validated test clips available\n")

    rows = []
    for p in sorted(glob.glob(str(RESULTS_DIR / "*_predictions.json"))):
        key = Path(p).name.replace("_predictions.json", "")
        try:
            preds = json.load(open(p, encoding="utf-8"))
        except Exception:
            continue
        hyp_by_id = {str(x.get("clip_id", "")): str(x.get("hypothesis",
                     x.get("prediction", ""))) for x in preds}
        sil_by_id = {str(x.get("clip_id", "")): str(x.get("reference", ""))
                     for x in preds}
        ids = [c for c in human if c in hyp_by_id]
        if not ids:
            continue
        h_refs = [human[c] for c in ids]
        m_hyps = [hyp_by_id[c] for c in ids]
        s_refs = [sil_by_id.get(c, "") for c in ids]
        rows.append({
            "key": key,
            "label": LABELS.get(key, key),
            "n": len(ids),
            "wer_human": _score(h_refs, m_hyps, "wer"),
            "cer_human": _score(h_refs, m_hyps, "cer"),
            "wer_silver": _score(s_refs, m_hyps, "wer"),
        })

    if not rows:
        print("[ERROR] No predictions JSON overlaps the validated clips.")
        raise SystemExit(1)

    rows.sort(key=lambda r: (r["wer_human"] if r["wer_human"] is not None else 9))
    print(f"{'Model':<34}{'n':>5}{'WER(human)':>12}{'CER(human)':>12}{'WER(silver)':>13}")
    print("-" * 76)
    for r in rows:
        print(f"{r['label']:<34}{r['n']:>5}{r['wer_human']!s:>12}"
              f"{r['cer_human']!s:>12}{r['wer_silver']!s:>13}")
    print("-" * 76)
    print("WER(human)  = model vs HUMAN transcript  <- gold-reference benchmark number")
    print("WER(silver) = model vs silver label on the same clips (for comparison)")

    out = RESULTS_DIR / "human_test_wer.json"
    json.dump(rows, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
