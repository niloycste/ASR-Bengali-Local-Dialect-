"""
Step 12: Build a stratified manual-validation sample for silver-standard quality.

The BanglaMix transcripts are silver-standard (Whisper-generated). To report
their human-verified accuracy, this script samples N clips (default 200) from a
split, stratified by dialect and code-switching status, and writes a CSV for a
human annotator to transcribe by listening. The filled CSV is then scored by
13_validation_wer.py to report silver-standard label WER/CER.

Output: evaluation/results/validation_sample.csv
  columns: clip_id, dialect, is_code_switched, audio_path,
           silver_transcript, human_transcript  (human_transcript is BLANK ---
           fill it by listening to each clip)

Usage:
    python 12_make_validation_sample.py                 # 200 clips from test
    python 12_make_validation_sample.py --n 300 --split test
"""
from __future__ import annotations

import argparse
import csv
import random
import sys
from collections import defaultdict
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(PROJECT_ROOT))

from fine_tuning.config import DATASET_DIR, RESULTS_DIR


def is_cs(row: dict) -> bool:
    return str(row.get("is_code_switched", "")).strip().lower() in ("true", "1")


def _write_xlsx(path: Path, cols: list[str], records: list[dict]) -> None:
    """Write an Excel .xlsx with Bengali-capable transcript fonts."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("[info] openpyxl not installed; skipping .xlsx (pip install openpyxl).")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "validation"
    ws.append(cols)

    header_font = Font(name="Nirmala UI", bold=True, color="FFFFFF")
    bengali_font = Font(name="Nirmala UI", size=11)
    green = PatternFill("solid", fgColor="2E7D32")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = green

    for rec in records:
        ws.append([rec[c] for c in cols])

    widths = {
        "clip_id": 22,
        "dialect": 12,
        "is_code_switched": 14,
        "audio_path": 60,
        "silver_transcript": 50,
        "human_transcript": 50,
    }
    for i, col in enumerate(cols, 1):
        ws.column_dimensions[chr(64 + i)].width = widths.get(col, 18)

    yellow = PatternFill("solid", fgColor="FFF9C4")
    for row in range(2, len(records) + 2):
        silver_cell = ws.cell(row, 5)
        silver_cell.font = bengali_font
        silver_cell.alignment = Alignment(wrap_text=True)

        human_cell = ws.cell(row, 6)
        human_cell.font = bengali_font
        human_cell.fill = yellow
        human_cell.alignment = Alignment(wrap_text=True)

    ws.freeze_panes = "A2"
    wb.save(path)


def _load_rows(split: str) -> list[dict]:
    """Load clips from the manifest; fall back to a predictions JSON."""
    csv_path = DATASET_DIR / split / "manifest.csv"
    if csv_path.exists():
        with open(csv_path, encoding="utf-8") as f:
            return [r for r in csv.DictReader(f) if str(r.get("transcript", "")).strip()]

    import glob
    import json

    candidates = [RESULTS_DIR / "ft_whisper_all_predictions.json"]
    candidates += [Path(p) for p in sorted(glob.glob(str(RESULTS_DIR / "*_predictions.json")))]
    for path in candidates:
        if path.exists():
            print(f"[info] manifest not found; using {path.name} as the clip source.")
            rows = []
            with open(path, encoding="utf-8") as f:
                predictions = json.load(f)
            for rec in predictions:
                if not str(rec.get("reference", "")).strip():
                    continue
                rows.append({
                    "clip_id": rec.get("clip_id"),
                    "audio_path": rec.get("audio_path", ""),
                    "dialect": rec.get("dialect", "unknown"),
                    "is_code_switched": rec.get("is_code_switched", False),
                    "transcript": rec.get("reference", ""),
                })
            return rows

    print(f"[ERROR] No manifest ({csv_path}) and no predictions in {RESULTS_DIR}.")
    print("        Run 05_build_dataset.py, or run the evaluation first.")
    raise SystemExit(1)


def main(n: int = 200, split: str = "test", seed: int = 42):
    rng = random.Random(seed)
    rows = _load_rows(split)

    by_dialect = defaultdict(list)
    for row in rows:
        by_dialect[row.get("dialect", "unknown")].append(row)

    dialects = sorted(by_dialect)
    per = max(1, n // len(dialects))

    sample, chosen = [], set()
    for dialect in dialects:
        clips = by_dialect[dialect]
        cs_clips = [clip for clip in clips if is_cs(clip)]
        bn_clips = [clip for clip in clips if not is_cs(clip)]
        rng.shuffle(cs_clips)
        rng.shuffle(bn_clips)

        half = per // 2
        pick = cs_clips[:half] + bn_clips[: per - len(cs_clips[:half])]
        for clip in pick[:per]:
            if clip["clip_id"] not in chosen:
                sample.append(clip)
                chosen.add(clip["clip_id"])

    rng.shuffle(rows)
    for row in rows:
        if len(sample) >= n:
            break
        if row["clip_id"] not in chosen:
            sample.append(row)
            chosen.add(row["clip_id"])

    sample = sample[:n]
    rng.shuffle(sample)

    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    cols = [
        "clip_id",
        "dialect",
        "is_code_switched",
        "audio_path",
        "silver_transcript",
        "human_transcript",
    ]
    records = [{
        "clip_id": row["clip_id"],
        "dialect": row.get("dialect", "unknown"),
        "is_code_switched": is_cs(row),
        "audio_path": row.get("audio_path", ""),
        "silver_transcript": str(row.get("transcript", "")).strip(),
        "human_transcript": "",
    } for row in sample]

    out = RESULTS_DIR / "validation_sample.csv"
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=cols)
        writer.writeheader()
        writer.writerows(records)

    xlsx = RESULTS_DIR / "validation_sample.xlsx"
    _write_xlsx(xlsx, cols, records)

    n_cs = sum(1 for row in sample if is_cs(row))
    print(f"Wrote {len(sample)} clips -> {xlsx}")
    print(f"                       -> {out}")
    print(
        f"  CS clips: {n_cs} | BN clips: {len(sample) - n_cs} | "
        f"dialects: {len({row.get('dialect') for row in sample})}"
    )
    print("\nNext: open validation_sample.xlsx, listen to each clip (audio_path), type")
    print("the true transcript into 'human_transcript', then run: python 13_validation_wer.py")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--n", type=int, default=200, help="number of clips (default 200)")
    parser.add_argument("--split", default="test", choices=["train", "dev", "test"])
    args = parser.parse_args()
    main(n=args.n, split=args.split)
